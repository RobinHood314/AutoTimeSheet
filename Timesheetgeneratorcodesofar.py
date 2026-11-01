#!/usr/bin/env python3

import datetime
from datetime import date
import openpyxl
import sys

book = openpyxl.load_workbook(filename='TimeSheet.xlsx', data_only=True)
timesheet = book.active

def getDate():
    today = datetime.datetime.now()
    thisYear = today.strftime('%Y')
    while True:
        dateComponents = input('Enter the Start Date in MM-DD-YYYY Format:').split('-')
        month, day, year = [int(item) for item in dateComponents]
        if day == 1 or day == 16 and year == int(thisyear):
            dateInCell = date(year, month, day)
            break
        else:
            print('Invalid date, please try again.')

    return dateInCell

def findEnd(month, day, year, leapYear):
    if day == 1:
        endCell = 21
    elif day == 16:
        if month == 2:
            if leapYear == True:
                endCell = 20
            else:
                endCell = 19
        elif month % 2 == 0 and month > 7 or month % 2 == 1 and month <= 7:
            endCell = 22
        else:
            endCell = 21
    else:
        dateInCell = getDate()
        if month == 2:
            if leapYear == True:
                endCell = 20
            else:
                endCell = 19
        elif month % 2 == 0 and month > 7 or month % 2 == 1 and month < 7:
            endCell = 22
        else:
            endCell = 21

    return endCell

def fillTop():
    print('Please enter your name:')
    timesheet['A2'].value = input()
    for columnNum in range(2,8):
        print('Enter scheduled hours for ' + str(timesheet.cell(row=1, column=columnNum).value + ':' ))
        timesheet.cell(row=2, column=columnNum).value = input() 

def verifyTop():
    print('Is the below information correct? (y for yes, n for no)')
    for columnNo in range(1,8):
        print(timesheet.cell(row=1, column=columnNo).value)
        print(timesheet.cell(row=2, column=columnNo).value)
        print('----------------------------')

    correct = input()
    while True:
        if correct == 'y':
            break
        elif correct == 'n':
            fillTop()
            break
        else:
            print('Please enter y or n.')

def fillDate():
    incDay = datetime.timedelta(days=1)

    if timesheet['A6'].value != None:
        if timesheet['A19'].value == None:
            timesheet['A6'].value = timesheet['A18'].value + incDay
        elif timesheet['A20'].value == None:
            timesheet['A6'].value = timesheet['A19'].value + incDay
        elif timesheet['A21'].value == None:
            timesheet['A6'].value = timesheet['A20'].value + incDay
        elif timesheet['A22'].value == None:
            timesheet['A6'].value = timesheet['A21'].value + incDay
        elif timesheet['A23'].value == None:
            timesheet['A6'].value = timesheet['A22'].value + incDay
        else:
            dateInCell = getDate()

        dateInCell = timesheet['A6'].value
    else:
        dateInCell = getDate()

    day = dateInCell.day
    month = dateInCell.month
    year = dateInCell.year

    if year % 4 == 0 and year % 100 != 0 or year % 400 == 0:
        leapYear = True
    else:
        leapYear = False
    
    endCell = findEnd(month, day, year, leapYear)

    for i in range(6,endCell):
        timesheet.cell(row = i, column = 1).value = dateInCell
        dateInCell += incDay

    for i in range(endCell,22):
        timesheet.cell(row = i, column = 1).value = None

    return endCell

def logHours():
    endCell = fillDate()
    for i in range(6, endCell):
        date = timesheet.cell(row = i, column = 1).value
        match date.strftime('%w'):
            case '0':
                print('Enter Hours Worked for', end=' ')
                print('Sunday', end=', ')
            case '1':
                print('Enter Hours Worked for', end=' ')
                print('Monday', end=', ')
            case '2':
                print('Enter Hours Worked for', end=' ')
                print('Tuesday', end=', ')
            case '3':
                print('Enter Hours Worked for', end=' ')
                print('Wednesday', end=', ')
            case '4':
                print('Enter Hours Worked for', end=' ')
                print('Thursday', end=', ')
            case '5':
                print('Enter Hours Worked for', end=' ')
                print('Friday', end=', ')
            case '6':
                continue
            case _:
                print('Error: Undefined Day')
            
        print(date.strftime('%B %d, %Y') + ':')
        for j in range(2, 10):
            if date.strftime('%w') == 6:
                continue

            if j < 8:
                if timesheet.cell(row=2, column=j).value == None:
                    continue
            print('How many hours of ' + timesheet.cell(row=5, column=j).value + '?')
            print('(Round to the nearest .25)')
            timesheet.cell(row=i, column=j).value = input()

        print('Any comments to add? (Leave blank if no)')
        timesheet.cell(row=i, column=10).value = input()

    return endCell

def verifyHours():
    endCell = logHours()
    print('\n \n \n')
    print('Hours worked this pay period:')
    print('\n')
    for i in range(6, endCell):
        date = timesheet.cell(row=i, column=1).value
        print(date.strftime('%B %d, %Y' + ':'))
        for j in range(2,12):
            if timesheet.cell(row=i, column=j).value == None:
                continue
            print('    ' + timesheet.cell(row=5, column=j).value + ': ' + str(timesheet.cell(row=i, column=j).value))
        
        print('\n')
    print('\n')
    print('Total Hours: ' + str(timesheet['K23'].value))


verifyTop()
verifyHours()
book.save('TimeSheet.xlsx')
